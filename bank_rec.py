"""Bank reconciliation: bank statement lines matched with the book lines of a bank account (511 / 512 / 519 / 53)."""
from __future__ import annotations

import json
from datetime import datetime
from decimal import Decimal

from database import display_date, iso_date, utcnow

ZERO = Decimal("0")


def migrate(db):
    db.execute("""CREATE TABLE IF NOT EXISTS bank_statement_lines (id INTEGER PRIMARY KEY, account_code TEXT NOT NULL, currency TEXT NOT NULL, line_date TEXT NOT NULL,
        description TEXT, reference TEXT, amount TEXT NOT NULL, journal_line_id INTEGER, created_at TEXT NOT NULL)""")


def _d(value):
    try: return Decimal(str(value if value not in (None, "") else 0).replace(",", ""))
    except Exception: raise ValueError(f"'{value}' is not an amount")


def add_statement_lines(database, account, currency, rows, user_id):
    """rows: [{date, description, reference, amount}] amount + = money in (deposit), - = money out."""
    account = str(account or "").split(" - ", 1)[0].strip(); currency = str(currency or "USD").upper()
    if not account.startswith(("511", "512", "519", "53")): raise ValueError("Choose a bank or cash account (511, 512, 519 or 53)")
    added = 0
    with database.connect() as db:
        for row in rows:
            amount = _d(row.get("amount"))
            if not amount: continue
            db.execute("INSERT INTO bank_statement_lines(account_code,currency,line_date,description,reference,amount,created_at) VALUES(?,?,?,?,?,?,?)",
                (account, currency, iso_date(row.get("date"), "Date"), str(row.get("description") or "").strip(), str(row.get("reference") or "").strip(), str(amount), utcnow())); added += 1
        db.execute("INSERT INTO audit_log(user_id,action,entity,details,created_at) VALUES(?,?,?,?,?)", (user_id, "import", "bank_statement", json.dumps({"account": account, "lines": added}), utcnow()))
    return added


def read_statement_file(path):
    """Excel / CSV bank statement: Date, Description, Reference and either Amount or Debit / Credit (or Withdrawal / Deposit)."""
    from importer import _header_map, _date
    if str(path).lower().endswith(".csv"):
        import csv
        with open(path, newline="", encoding="utf-8-sig") as handle: table = [row for row in csv.reader(handle)]
        class Sheet:
            def iter_rows(self, min_row=1, max_row=None, values_only=True):
                return [tuple(r) for r in table[min_row - 1:max_row]]
        sheet = Sheet(); close = lambda: None
    else:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True); sheet = book.worksheets[0]; close = book.close
    try:
        header, columns = _header_map(sheet, {"date": ("date", "تاريخ", "value date"), "description": ("description", "details", "narration", "بيان", "libell"),
            "reference": ("reference", "ref", "cheque", "check"), "amount": ("amount", "montant", "مبلغ"), "debit": ("debit", "withdrawal", "out", "مدين"),
            "credit": ("credit", "deposit", "in ", "دائن")})
        rows = []
        for row in sheet.iter_rows(min_row=header + 1, values_only=True):
            get = lambda f: row[columns[f]] if f in columns and columns[f] < len(row) else None
            if not any(v not in (None, "") for v in row) or get("date") in (None, ""): continue
            if "amount" in columns: amount = _d(get("amount"))
            else: amount = _d(get("credit")) - _d(get("debit"))  # bank statement: credit = money in, debit = money out
            rows.append({"date": _date(get("date")), "description": get("description") or "", "reference": get("reference") or "", "amount": amount})
        return rows
    finally: close()


def statement_lines(database, account, date_from, date_to):
    with database.connect() as db:
        return [dict(r) for r in db.execute("SELECT * FROM bank_statement_lines WHERE account_code=? AND line_date BETWEEN ? AND ? ORDER BY line_date,id", (account, date_from, date_to))]


def book_lines(database, account, currency, date_from, date_to):
    """Book movements of the bank account: debit = money in (+), credit = money out (-)."""
    with database.connect() as db:
        matched = {r["journal_line_id"]: r["id"] for r in db.execute("SELECT id,journal_line_id FROM bank_statement_lines WHERE journal_line_id IS NOT NULL")}
        rows = [dict(r) for r in db.execute("""SELECT l.id,e.entry_number,e.entry_date,COALESCE(l.description,e.description) description,e.description entry_description,
            CAST(l.debit AS REAL) debit,CAST(l.credit AS REAL) credit,COALESCE(NULLIF(l.line_currency,''),e.currency) currency,
            CASE WHEN COALESCE(NULLIF(l.line_currency,''),e.currency)<>e.currency AND l.amount IS NOT NULL THEN CAST(l.amount AS REAL) ELSE NULL END line_amount
            FROM journal_lines l JOIN journal_entries e ON e.id=l.entry_id JOIN accounts a ON a.id=l.account_id WHERE a.code=?""", (account,))]
    result = []
    for row in rows:
        try: day = iso_date(row["entry_date"])
        except ValueError: continue
        if not date_from <= day <= date_to or row["currency"] != currency: continue
        sign = 1 if row["debit"] else -1
        value = Decimal(str(row["line_amount"])) * sign if row["line_amount"] is not None else _d(row["debit"]) - _d(row["credit"])
        result.append({**row, "iso_date": day, "amount": value, "statement_id": matched.get(row["id"])})
    return sorted(result, key=lambda r: (r["iso_date"], r["id"]))


def auto_match(database, account, currency, date_from, date_to, days=5):
    """One-to-one: same amount, dates within `days` days, the closest date first."""
    statement = [s for s in statement_lines(database, account, date_from, date_to) if not s["journal_line_id"] and s["currency"] == currency]
    books = [b for b in book_lines(database, account, currency, date_from, date_to) if not b["statement_id"]]
    used = set(); matched = 0
    with database.connect() as db:
        for line in statement:
            day = datetime.strptime(line["line_date"], "%Y-%m-%d"); amount = _d(line["amount"])
            candidates = [b for b in books if b["id"] not in used and abs(b["amount"] - amount) < Decimal("0.005")
                          and abs((datetime.strptime(b["iso_date"], "%Y-%m-%d") - day).days) <= days]
            if not candidates: continue
            best = min(candidates, key=lambda b: abs((datetime.strptime(b["iso_date"], "%Y-%m-%d") - day).days))
            used.add(best["id"]); matched += 1
            db.execute("UPDATE bank_statement_lines SET journal_line_id=? WHERE id=?", (best["id"], line["id"]))
    return matched


def match(database, statement_id, journal_line_id):
    with database.connect() as db:
        line = db.execute("SELECT * FROM bank_statement_lines WHERE id=?", (int(statement_id),)).fetchone()
        if not line: raise KeyError("Statement line not found")
        if db.execute("SELECT 1 FROM bank_statement_lines WHERE journal_line_id=? AND id<>?", (int(journal_line_id), int(statement_id))).fetchone():
            raise ValueError("This book line is already matched with another statement line")
        db.execute("UPDATE bank_statement_lines SET journal_line_id=? WHERE id=?", (int(journal_line_id), int(statement_id)))
    return True


def unmatch(database, statement_id):
    with database.connect() as db: db.execute("UPDATE bank_statement_lines SET journal_line_id=NULL WHERE id=?", (int(statement_id),))
    return True


def delete_statement_line(database, statement_id):
    with database.connect() as db: db.execute("DELETE FROM bank_statement_lines WHERE id=?", (int(statement_id),))
    return True


def post_statement_line(database, statement_id, account_code, user_id):
    """Bank charges, interest, transfers seen only on the statement: book them and match them."""
    with database.connect() as db: line = db.execute("SELECT * FROM bank_statement_lines WHERE id=?", (int(statement_id),)).fetchone()
    if not line: raise KeyError("Statement line not found")
    if line["journal_line_id"]: raise ValueError("This statement line is already matched")
    amount = _d(line["amount"]); other = str(account_code or "").split(" - ", 1)[0].strip()
    if not other: raise ValueError("Choose the account (for example 6278 bank charges)")
    bank_side, other_side = ("D", "C") if amount > 0 else ("C", "D")
    voucher = database.save_journal_voucher({"entry_date": display_date(line["line_date"]), "description": f"Bank: {line['description'] or 'statement line'}", "currency": line["currency"], "voucher_type": "03"},
        [{"account_code": line["account_code"], "line_currency": line["currency"], "side": bank_side, "amount": str(abs(amount)), "reference": line["reference"], "description": line["description"]},
         {"account_code": other, "line_currency": line["currency"], "side": other_side, "amount": str(abs(amount)), "reference": line["reference"], "description": line["description"]}], user_id)
    with database.connect() as db:
        bank_line = db.execute("""SELECT l.id FROM journal_lines l JOIN accounts a ON a.id=l.account_id JOIN journal_entries e ON e.id=l.entry_id
            WHERE e.entry_number=? AND a.code=? ORDER BY l.id LIMIT 1""", (voucher["voucher"]["entry_number"], line["account_code"])).fetchone()
        db.execute("UPDATE bank_statement_lines SET journal_line_id=? WHERE id=?", (bank_line["id"], int(statement_id)))
    return voucher["voucher"]["entry_number"]


def reconciliation(database, account, currency, date_from, date_to, statement_balance=None):
    """Bank balance per statement = book balance + payments not yet cleared - deposits not yet cleared + statement items not booked."""
    books_all = book_lines(database, account, currency, "0000-01-01", date_to); books = [b for b in books_all if b["iso_date"] >= date_from]
    statement = statement_lines(database, account, date_from, date_to)
    book_balance = sum((b["amount"] for b in books_all), ZERO)
    outstanding = [b for b in books if not b["statement_id"]]
    not_booked = [s for s in statement if not s["journal_line_id"] and s["currency"] == currency]
    deposits = sum((b["amount"] for b in outstanding if b["amount"] > 0), ZERO); payments = -sum((b["amount"] for b in outstanding if b["amount"] < 0), ZERO)
    bank_only = sum((_d(s["amount"]) for s in not_booked), ZERO)
    expected = book_balance - deposits + payments + bank_only
    difference = (_d(statement_balance) - expected) if statement_balance not in (None, "") else None
    rows = [["Balance per books at " + display_date(date_to), book_balance.quantize(Decimal("0.01"))],
            ["Less: deposits in the books not yet on the statement", (-deposits).quantize(Decimal("0.01"))],
            ["Add: payments in the books not yet cleared by the bank", payments.quantize(Decimal("0.01"))],
            ["Add / less: statement items not yet booked (charges, interest...)", bank_only.quantize(Decimal("0.01"))],
            ["EXPECTED BALANCE PER BANK STATEMENT", expected.quantize(Decimal("0.01"))]]
    if difference is not None: rows += [["Balance per bank statement (entered)", _d(statement_balance).quantize(Decimal("0.01"))], ["DIFFERENCE (should be 0)", difference.quantize(Decimal("0.01"))]]
    sections = [{"heading": f"Bank reconciliation - account {account} ({currency})", "headers": ["Item", f"Amount ({currency})"], "rows": rows, "total_rows": [4] + ([6] if difference is not None else [])},
                {"heading": "Book items not yet on the statement", "headers": ["Date", "Voucher", "Description", "Amount"],
                 "rows": [[display_date(b["iso_date"]), b["entry_number"], b["description"] or "", b["amount"].quantize(Decimal("0.01"))] for b in outstanding] or [["None", "", "", ""]], "total_rows": []},
                {"heading": "Statement items not yet booked", "headers": ["Date", "Reference", "Description", "Amount"],
                 "rows": [[display_date(s["line_date"]), s["reference"] or "", s["description"] or "", _d(s["amount"]).quantize(Decimal("0.01"))] for s in not_booked] or [["None", "", "", ""]], "total_rows": []}]
    company = database.settings()
    return {"title": "Bank Reconciliation", "meta": [f"Company: {company.get('company_name') or '-'}   Account {account}   Period {display_date(date_from)} to {display_date(date_to)}"],
            "sections": sections, "difference": float(difference) if difference is not None else None, "expected": float(expected), "book_balance": float(book_balance),
            "matched": sum(1 for s in statement if s["journal_line_id"]), "statement_count": len(statement)}
